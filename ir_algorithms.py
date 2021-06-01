import openpyxl # for reading excel files
import re, regex, pickle, numpy as np
import heapq

frequent_terms_num = 70 # removing # of most frequent terms from dictionary

arabic_persian_chars = [['ي', 'ی'], ['ئ', 'ی'], ['ك', 'ک'], ['ة', 'ه'], ['ؤ', 'و'],\
             ['آ', 'ا'], ['إ', 'ا'], ['أ', 'ا'], ['ٱ', 'ا'], ['ء', '']]


class Document:
    def __init__(self, doc_id, content, url) -> None:
        self.doc_id = doc_id
        self.content = content
        self.url = url


class Posting:
    def __init__(self, doc_id, freq) -> None:
        self.doc_id = doc_id
        self.freq = freq

    def __str__(self) -> str:
        return 'doc_id: ' + str(self.doc_id) + '\tfreq: ' + str(self.freq)

    def __repr__(self) -> str:
        return str(self)


class PostingsList:
    def __init__(self) -> None:
        self.plist = []
        self.term_freq = 0

    def __str__(self) -> str:
        return 'term_freq: ' + str(self.term_freq) + '\t' + str(self.plist)

    def __repr__(self) -> str:
        return str(self)

    def __lt__(self, other):
        return self.term_freq < other.term_freq

    
    


class IR:
    def __init__(self) -> None:
        self.dictionary = dict()
        self.documents = None
        self.docs_dict = dict()

    
    # building the inverted index
    def build_inverted_index(self, file_name):
        self.init_file(file_name)
        for doc in self.documents:
            self.index_document(doc)
        print('Inverted Index Matrix construction completed')
        self.build_docs_dict()
        # removing most frequent items
        self.remove_frequents(frequent_terms_num)
        # saving the dictionary
        with open('index.pickle', 'wb') as handle:
            pickle.dump(self.dictionary, handle, protocol=pickle.HIGHEST_PROTOCOL)
        with open('docs_dict.pickle', 'wb') as handle:
            pickle.dump(self.docs_dict, handle, protocol=pickle.HIGHEST_PROTOCOL)


    # building a dictionary mapping documents IDs to URL
    def build_docs_dict(self):
        for doc in self.documents:
            self.docs_dict[doc.doc_id] = doc.url


    
    # loading the existing inverted index from file
    def load_inverted_index(self):
        with open('index.pickle', 'rb') as handle:
            self.dictionary = pickle.load(handle)
        with open('docs_dict.pickle', 'rb') as handle:
            self.docs_dict = pickle.load(handle)



    # initializing documents list by reading the excel dataset
    def init_file(self, file_name):
        wb_obj = openpyxl.load_workbook(file_name)
        sheet = wb_obj.active
        headers = []
        dataset = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i == 0:
                for header in row:
                    headers.append(header)
            else:
                document = Document(row[0], row[1], row[2])
                dataset.append(document)
        self.documents = dataset
        print('Initialized Excel file')


    # processing the documents one by one for building the index
    def index_document(self, doc):
        tokens = self.get_tokens(doc.content)
        counts = self.get_counts_dict(tokens)
        unique_tokens = counts.keys()
        for unique_token in unique_tokens:
            posting = Posting(doc.doc_id, counts[unique_token])
            self.add_posting(posting, unique_token)

    
    # get tokens for each document
    def get_tokens(self, text):
        '''
        tokens = re.split('!|,|[|]|\{|\}|\s|-|_|\(|\)|\.|؟|:|»|«|\(|\)|؛|،|\*|&|\
            \^|%|\$|#|@|~|\\|\"|"|\'|;|>|<|\||=|\+|\?', text)
        tokens = list(filter(None, tokens))
        '''
        tokens = regex.findall(r'[\p{Cf}\p{L}]+', text)
        # modifying tokens
        tokens = [self.modify_token(token) for token in tokens if self.modify_token(token) != '']
        return tokens

    
    # getting a dictionary of unique terms and the frequencies of each term in a list
    def get_counts_dict(self, tokens):
        counts = dict()
        for token in tokens:
            if token not in counts:
                counts[token] = 1
            else:
                counts[token] += 1
        return counts


    # add posting to the postings_list of the corresponding term in dictionary
    def add_posting(self, posting, term):
        postings_list = None
        if term not in self.dictionary:
            postings_list = PostingsList()
        else:
            postings_list = self.dictionary[term]
        postings_list.term_freq += 1
        postings_list.plist.append(posting)
        self.dictionary[term] = postings_list
        
            

    
    # modify token with stemming, tokenization, normalization, etc
    def modify_token(self, token):
        # to be implemented...
        #
        #
        #
        # m = re.findall(r"^ب[ا-ی]*ید$", token)
        # if m:
        #     if m[0] not in unique_list:
        #         unique_list.append(m[0])
        # removing stop words
        #if len(token) < 3:
        #    return ''
        result = self.normalize(token)
        return result


    # normalize the given token
    def normalize(self, token):
        # normalizing characters
        result = token
        # list of [Arabic_character, Persian_character]
        global arabic_persian_chars
        for i in range(len(arabic_persian_chars)):
            if arabic_persian_chars[i][0] in token:
                result = result.replace(arabic_persian_chars[i][0], arabic_persian_chars[i][1])
        '''
        for char_set in arabic_persian_chars:
            result = re.sub(char_set[0], char_set[1], result)
        '''
        if not token.isalnum() and '\u200c' not in token and '\xad' not in token:
            print(token)
        return result


    # processing queries
    def process_query(self, query):
        tokens = self.get_tokens(query)
        if len(tokens) == 1:
            self.process_query_single_word(tokens[0])
        elif len(tokens) > 1:
            self.process_query_mult_words(tokens)
        else:
            print('Query is not valid; please make sure your query contains words.')
    

    # processing queries with only one word
    def process_query_single_word(self, query):
        result_ids = self.get_posting_ids(query)
        if len(result_ids) == 0:
            print("No result found!")
        else:
            for index in result_ids:
                print(index, self.docs_dict[index])
        return result_ids


    # processing queries with multiple words - alternative solution using intersection
    def process_query_mult_words_alt(self, tokens):
        id_lists = []
        pointers = []
        result_set = []
        for t in tokens:
            posting_ids = self.get_posting_ids(t)
            id_lists.append(posting_ids)
            if len(posting_ids) > 0:
                pointers.append(0)
            else:
                pointers.append(None)
        terminated = all(x is None for x in pointers)
        while not terminated:
            min_index = len(self.documents) + 1
            min_pointer_ind = len(pointers)
            for i in range(len(pointers)):
                if pointers[i] is None:
                    continue
                id_list = id_lists[i]
                pointer = pointers[i]
                if id_list[pointer] < min_index:
                    min_index = id_list[pointer]
                    min_pointer_ind = i
            # found minimum
            if result_set == [] or result_set[len(result_set)-1][0] != min_index:
                result_set.append([min_index, 1])
            else:
                result_set[len(result_set)-1][1] += 1
            # move min_pointer one step further
            pointers[min_pointer_ind] += 1
            if pointers[min_pointer_ind] >= len(id_lists[min_pointer_ind]):
                pointers[min_pointer_ind] = None
            terminated = all(x is None for x in pointers)
        result_set = sorted(result_set, key=lambda item: item[1], reverse=True)
        return result_set



    # processing queries with multiple words
    def process_query_mult_words(self, tokens):
        id_lists = []
        result_set = dict()
        for t in tokens:
            posting_ids = self.get_posting_ids(t)
            id_lists.append(posting_ids)
        for id_list in id_lists:
            for item in id_list:
                if item not in result_set:
                    result_set[item] = 1
                else:
                    result_set[item] += 1
        result_set = sorted(result_set.items(), key=lambda item: (-item[1], item[0]))
        last_count = 0
        for item in result_set:
            if last_count != item[1]:
                print('\nNumber of Words:', item[1])
                last_count = item[1]
            index = item[0]
            print(index, self.docs_dict[index])
        return result_set



    # getting document IDs for a given term
    def get_posting_ids(self, term):
        ids = []
        if term in self.dictionary:
            postings_list = self.dictionary[term]
            postings = postings_list.plist
            for p in postings:
                ids.append(p.doc_id)
        return ids


    # removing k most frequent term from dictionary
    def remove_frequents(self, k):
        most_frequent = heapq.nlargest(k, self.dictionary, key=self.dictionary.get)
        for key in most_frequent:
            del self.dictionary[key]

        